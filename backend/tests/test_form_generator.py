"""
Tests for compliance/form_generator.py — see its module docstring for why
this only ever surfaces facts already in FORM_CATALOG (compliance/
form_navigator.py), never invented fees or fields.

Usage:
    python -m pytest tests/test_form_generator.py -v
"""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from compliance.form_generator import UnknownFormId, generate_form_docx
from compliance.form_navigator import FORM_CATALOG


def test_unknown_form_id_raises():
    with pytest.raises(UnknownFormId):
        generate_form_docx("NOT_A_REAL_FORM_ID")


@pytest.mark.parametrize("form", FORM_CATALOG, ids=lambda f: f["form_id"])
def test_every_catalog_form_generates_a_valid_docx(form):
    """Every real form in the catalog must actually generate — a
    regression here means a form a user is told exists (via
    actionable_forms) can't actually be downloaded."""
    docx_bytes = generate_form_docx(form["form_id"])
    assert docx_bytes[:2] == b"PK"  # .docx is a zip archive

    doc = Document(BytesIO(docx_bytes))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert form["title"] in full_text
    assert form["statutory_mandate"] in full_text
    assert form["submission_portal"] in full_text
    assert form["deadline"] in full_text
    for attachment in form["required_attachments"]:
        assert attachment in full_text


def test_generated_docx_names_this_a_cover_sheet_not_the_official_form():
    """The generated document must be honest about what it is — see this
    module's docstring on why that framing matters."""
    docx_bytes = generate_form_docx(FORM_CATALOG[0]["form_id"])
    doc = Document(BytesIO(docx_bytes))
    full_text = "\n".join(p.text for p in doc.paragraphs).lower()
    assert "not the official government form" in full_text
