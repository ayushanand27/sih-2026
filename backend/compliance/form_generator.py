"""
Fillable-form generator for IP-SAKTI.

Turns one compliance/form_navigator.py FORM_CATALOG entry into a
downloadable .docx cover sheet — the form's title, statutory mandate,
submission portal, required attachments, and deadline, exactly as that
already-verified catalog states them, plus blank fields for the applicant
to fill in by hand (name, address, contact, date, signature).

Deliberately does NOT invent content the catalog doesn't have: no fee
amount, no processing-time estimate beyond what's already in `deadline`,
no field list beyond `required_attachments`. This is a checklist/cover
sheet to prepare with and file alongside the real official form
downloaded from `submission_portal` — never a claim to *be* the official
government form, and the generated document says so explicitly.

Usage:
    python -m compliance.form_generator NBA_FORM_7 output.docx
"""

from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from compliance.form_navigator import FORM_CATALOG


class UnknownFormId(ValueError):
    """Raised when form_id isn't in FORM_CATALOG."""


def _find_form(form_id: str) -> dict:
    for form in FORM_CATALOG:
        if form["form_id"] == form_id:
            return form
    raise UnknownFormId(
        f"No form with id {form_id!r} in compliance.form_navigator.FORM_CATALOG."
    )


def generate_form_docx(form_id: str) -> bytes:
    """
    Build a fillable cover-sheet/checklist .docx for `form_id`. Every fact
    in it (title, statutory mandate, submission portal, required
    attachments, deadline) comes directly from FORM_CATALOG — nothing here
    is generated, estimated, or guessed.
    """
    form = _find_form(form_id)
    doc = Document()

    heading = doc.add_heading(form["title"], level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph(f"{form['agency']} — {form['jurisdiction'].title()}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].italic = True

    doc.add_paragraph()

    disclaimer_run = doc.add_paragraph().add_run(
        "This is a preparation cover sheet generated from IP-SAKTI "
        "Sahayak's form catalog — not the official government form "
        "itself. Download and file the actual form at the submission "
        "portal below, and verify current field requirements and fees "
        "directly with the issuing authority, since those can change "
        "independently of this catalog."
    )
    disclaimer_run.italic = True
    disclaimer_run.font.size = Pt(9)
    disclaimer_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.add_heading("Statutory mandate", level=2)
    doc.add_paragraph(form["statutory_mandate"])

    doc.add_heading("Submission portal", level=2)
    doc.add_paragraph(form["submission_portal"])

    doc.add_heading("Deadline", level=2)
    doc.add_paragraph(form["deadline"])

    doc.add_heading("Required attachments", level=2)
    for item in form["required_attachments"]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Applicant details (fill in)", level=2)
    for label in ("Name / organization", "Address", "Contact (phone / email)", "Date"):
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run("_" * 50)

    doc.add_paragraph()
    signature = doc.add_paragraph()
    signature.add_run("Signature: ").bold = True
    signature.add_run("_" * 40)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python -m compliance.form_generator <form_id> <output.docx>")
        sys.exit(1)
    form_id_arg, output_path = sys.argv[1], sys.argv[2]
    try:
        docx_bytes = generate_form_docx(form_id_arg)
    except UnknownFormId as exc:
        print(str(exc))
        sys.exit(1)

    # CLI args are an untrusted source (CWE-22): resolve and confirm the
    # target stays under the current directory before writing, instead of
    # passing an arbitrary caller-supplied path straight to open().
    cwd = Path.cwd().resolve()
    resolved_path = (cwd / output_path).resolve()
    if resolved_path != cwd and cwd not in resolved_path.parents:
        print(f"Refusing to write outside the current directory: {output_path}")
        sys.exit(1)

    with open(resolved_path, "wb") as f:
        f.write(docx_bytes)
    print(f"Wrote {resolved_path}")
